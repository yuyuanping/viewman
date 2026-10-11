"""磁力链接登记表增量维护：新出现的 magnet 链接与 xlsx 主表按 btih 比对，缺则补一行。

机制：不挂在消息处理路径上，而是轮询 `messages` 表里含 magnet 的行。班车是 sweeper
的定时器，但 bot 落库一条含 magnet 的消息时会立即唤醒同一轮同步，所以「收到即登记」
而不用等一个周期。（合并转发的正文是晚于行插入才回填的，所以游标只在写表成功后推进；
回填过的行由 `dirty` 点名重扫——它的 id 可能已经低于游标。写表前还会用表里的现值复查
一次 btih，抢跑或人工重复都不会再追加出行。）

解析规则不在这里重复实现：动态加载 data_dir 下的 extract_magnets.py，与离线全量重扫
共用同一套正则与截断规则，避免两份逻辑漂移。

列位置一律按表头文字现取，不硬编码下标——主表列数会随 extract_magnets.py 的版本变，
用户手工增删列也很常见。
"""
import asyncio
import importlib.util
import logging
import re
import time
from pathlib import Path

logger = logging.getLogger("saver.magnets")

SHEET = "磁力链接(去重)"
BATCH = 500          # 单次最多扫描的 magnet 行数

# 表头关键词 -> 本模块的规范字段名（按子串匹配，容忍「首次/最后」这类前缀）
_HEADER_KEYS = (
    ("序号", "seq"),
    ("特征码", "hash"),
    ("资源名称", "dn"),
    ("标题", "title"),
    ("大小", "size"),
    ("出现次数", "count"),
    ("首次日期", "first_date"),
    ("首次时间", "first_time"),
    ("最后日期", "last_date"),
    ("最后时间", "last_time"),
    ("发送者", "sender"),
    ("群号", "gid"),
    ("磁力链接", "link"),
    ("备注", "note"),
)


def _load_rules(path: Path):
    spec = importlib.util.spec_from_file_location("magnet_rules", str(path))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    for name in ("MAGNET_RAW", "parse_magnet", "usable_title", "SIZE_RE"):
        if not hasattr(mod, name):
            raise AttributeError(f"{path} 缺少 {name}")
    return mod


class MagnetRegistry:
    def __init__(self, xlsx: Path, rules):
        self.xlsx = Path(xlsx)
        self.rules = rules
        self.seen: set[str] = set()
        self.seq = 0           # 表里已用的最大「序号」，下一条 = self.seq + 1
        self.cursor = 0        # 已确认落表的 messages.id 水位
        self._scan_max = 0     # 本轮扫到的最大 id，flush 成功后才提交为 cursor
        self.dirty: set[int] = set()   # 文本被回填过的 messages.id
        self._dirty_batch: list[int] = []
        self.buffer: list[dict] = []
        self._err = ""
        self.disabled = False    # 表头/表缺失时停用，不影响消息存档

    @classmethod
    def build(cls, config: dict):
        data_dir = Path(config["data_dir"])
        xlsx = Path(config.get("magnet_xlsx") or data_dir / "磁力链接汇总.xlsx")
        script = Path(config.get("magnet_rules") or data_dir / "extract_magnets.py")
        if not xlsx.is_file():
            logger.warning("[MAGNET] 登记表不存在，跳过登记: %s", xlsx)
            return None
        if not script.is_file():
            logger.error("[MAGNET] 解析规则脚本缺失，跳过登记: %s", script)
            return None
        try:
            return cls(xlsx, _load_rules(script))
        except Exception as e:
            logger.error("[MAGNET] 初始化失败，跳过登记: %s: %s", type(e).__name__, e)
            return None

    # ---- 读表 ----
    def _read_sheet(self, wb) -> dict:
        """现值：列位置、已有 btih、已用序号、行数。表缺失/表头认不出会抛。"""
        ws = wb[SHEET]
        hdr = [str(v or "") for v in next(ws.iter_rows(max_row=1, values_only=True))]
        cols = self._cols(hdr)
        if "hash" not in cols:
            raise KeyError(f"表头没有『特征码』列: {hdr}")
        h_idx, s_idx = cols["hash"], cols.get("seq")
        hashes: set[str] = set()
        seqs: set[int] = set()
        nrows = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            nrows += 1
            if len(row) > h_idx:
                h = str(row[h_idx] or "").strip().lower()
                if h:
                    hashes.add(h)
            if s_idx is not None and len(row) > s_idx:
                try:
                    seqs.add(int(float(row[s_idx])))
                except (TypeError, ValueError):
                    pass
        return {"cols": cols, "hashes": hashes, "seqs": seqs, "nrows": nrows}

    @staticmethod
    def _base_seq(seqs: set[int], nrows: int) -> int:
        """已用最大「序号」；整列都没数字时退回行数，别从 1 开始和老行撞号。"""
        return max(seqs) if seqs else nrows

    def _load(self):
        import openpyxl
        # data_only：「特征码」列若是公式，取缓存值而不是 "=..." 文本
        wb = openpyxl.load_workbook(self.xlsx, read_only=True, data_only=True)
        try:
            state = self._read_sheet(wb)
        finally:
            wb.close()
        self.seen = state["hashes"]
        self.seq = self._base_seq(state["seqs"], state["nrows"])

    @staticmethod
    def _cols(header: list[str]) -> dict:
        """表头文字 -> 列下标；同一字段只认第一次出现。"""
        cols = {}
        for idx, name in enumerate(header):
            for key, field in _HEADER_KEYS:
                if key in name and field not in cols:
                    cols[field] = idx
                    break
        return cols

    async def load(self) -> None:
        # 不经 db._run：读表不碰数据库，别为了它占着 DB 锁
        try:
            await asyncio.to_thread(self._load)
        except Exception as e:
            self.disabled = True
            logger.error("[MAGNET] 登记表不可用，本次运行不再登记: %s: %s",
                         type(e).__name__, e)
            return
        logger.info("[MAGNET] 登记表就绪: %d 个特征码, 序号从 %d 起",
                    len(self.seen), self.seq + 1)

    # ---- 抽取 ----
    def _row_for(self, text, scope, peer_id, nick, ts, fm) -> dict | None:
        link, h, kind, dn, note = self.rules.parse_magnet(fm.group(0))
        if not link or not h or h in self.seen:
            return None
        self.seen.add(h)
        text = text or ""
        # 上下文按「行」取：合并转发的正文是多行渲染，整段拼过来会把
        # 上一个发言人的话当成标题
        line_start = text.rfind("\n", 0, fm.start()) + 1
        nl = text.find("\n", fm.end())
        line_end = len(text) if nl < 0 else nl
        head = re.sub(r"\s+", " ", text[line_start:fm.start()]).strip(" ：:　\t")
        tail = text[fm.end():line_end]
        next_line = text[line_end + 1:line_end + 201]
        when = time.localtime(ts or time.time())
        d, t = time.strftime("%Y-%m-%d", when), time.strftime("%H:%M:%S", when)
        who, when2 = str(nick or ""), t
        strip = getattr(self.rules, "strip_prefix", None)
        if strip:
            speaker, body, hhmmss = strip(head)
            if speaker and hhmmss:      # 合并转发内层：发言人/时间以内层为准
                who, when2 = speaker, hhmmss
                head = body
        title = head if self.rules.usable_title(head) else ""
        ms = (self.rules.SIZE_RE.search(tail)
              or self.rules.SIZE_RE.search(next_line)
              or self.rules.SIZE_RE.search(head))
        return dict(
            hash=h, dn=dn, title=title[:300],
            size=f"{ms.group(1)} {ms.group(2)}" if ms else "",
            count=1, first_date=d, first_time=when2, last_date=d, last_time=when2,
            sender=who, gid=str(peer_id or "") if scope == "group" else "",
            link=link, note=("实时新增" + ("；" + note if note else ""))[:200],
        )

    def _rows_for(self, text, scope, peer_id, nick, ts, mid=None) -> list[dict]:
        out = []
        for fm in self.rules.MAGNET_RAW.finditer(text or ""):
            try:
                rec = self._row_for(text, scope, peer_id, nick, ts, fm)
            except Exception as e:
                logger.warning("[MAGNET] 解析失败 msg.id=%s: %s: %s", mid,
                               type(e).__name__, e)
                continue
            if rec:
                out.append(rec)
        return out

    # ---- 写表 ----
    def _flush(self) -> int:
        """追加 buffer 里表里还没有的行。返回被复查挡掉的行数。"""
        import openpyxl
        # seen 只在启动时读过一次；期间用户可能手工删过行、离线脚本可能重写过表，
        # 也可能上一轮 sweeper 抢在 load 之前跑过 —— 都以表里的现值为准。
        wb_r = openpyxl.load_workbook(self.xlsx, read_only=True, data_only=True)
        try:
            state = self._read_sheet(wb_r)
        finally:
            wb_r.close()
        cols = state["cols"]
        keep = [r for r in self.buffer if r["hash"] not in state["hashes"]]
        dropped = len(self.buffer) - len(keep)
        if not keep:
            return dropped
        # 这一次不能带 data_only：openpyxl 会把公式换成缓存值，保存即毁掉公式
        wb = openpyxl.load_workbook(self.xlsx)
        try:
            ws = wb[SHEET]
            width = max(cols.values()) + 1
            taken = set(state["seqs"])
            for rec in keep:
                rec["seq"] = self._base_seq(taken, state["nrows"]) + 1
                taken.add(rec["seq"])
                row = [None] * width
                for field, idx in cols.items():
                    if field in rec:
                        row[idx] = rec[field]
                ws.append(row)
            wb.save(self.xlsx)
        finally:
            wb.close()
        self.seen |= state["hashes"] | {r["hash"] for r in keep}
        self.seq = self._base_seq(taken, state["nrows"])
        return dropped

    def _commit(self) -> None:
        self.cursor = max(self.cursor, self._scan_max)
        self.dirty.difference_update(self._dirty_batch)
        self._dirty_batch = []

    async def sync(self, db) -> None:
        if self.disabled:
            return
        if not self.buffer:
            # dirty：文本被回填过的行，id 可能已经低于游标，不点名就永远扫不到
            self._dirty_batch = sorted(self.dirty)[:BATCH]
            where = "id>?"
            args = [self.cursor]
            if self._dirty_batch:
                where += " OR id IN (%s)" % ",".join("?" * len(self._dirty_batch))
                args += self._dirty_batch
            rows = await db._run(lambda c: c.execute(
                f"SELECT id, scope, peer_id, sender_nick, ts, text FROM messages "
                f"WHERE ({where}) AND text LIKE '%magnet:%' ORDER BY id LIMIT ?",
                (*args, BATCH)).fetchall())
            rows = [dict(r) for r in rows]
            for r in rows:
                self._scan_max = max(self._scan_max, r["id"])
                new = self._rows_for(r["text"], r["scope"], r["peer_id"],
                                     r["sender_nick"], r["ts"], r["id"])
                if new:
                    self.buffer.extend(new)
                    logger.info("[MAGNET] 待登记 %s",
                                ", ".join(x["hash"][:12] for x in new))
        if not self.buffer:
            self._commit()
            return
        try:
            n = len(self.buffer)
            dropped = await asyncio.to_thread(self._flush)
            self._commit()
            self.buffer.clear()
            self._err = ""
            logger.info("[MAGNET] 已写入 %d 行%s -> %s", n - dropped,
                        f"（复查挡掉 {dropped} 行：表里已有）" if dropped else "",
                        self.xlsx)
        except Exception as e:
            # Excel 开着这个文件时写不进去（PermissionError）：攒着，下一轮 sweeper
            # 再试，一条不丢
            if not self._err:
                logger.error("[MAGNET] 写入失败，%d 行待重试: %s: %s",
                             len(self.buffer), type(e).__name__, e)
            self._err = str(e)
