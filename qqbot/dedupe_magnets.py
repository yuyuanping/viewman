"""磁力登记表去重：磁力链接(去重) 按特征码只留首次出现的那一行。

背景：早期 bot 版本在登记表晚于 sweeper 载入时会把同一 btih 按 seq=1,2
重复写多行（bot.py 里「曾在 xlsx 里堆重复行」的注释即此事），存量表里有
残留。现行代码写表前已按表内现值复查，不会再堆；本脚本只负责清历史。

用法：python dedupe_magnets.py
表被 Excel/WPS 打开时写不进去（PermissionError），关掉后重跑即可。
"""
import json
import os
import sys
import time
from pathlib import Path

_CFG = Path(__file__).resolve().parent / "config.json"
_DATA_DIR = (json.loads(_CFG.read_text(encoding="utf-8")).get("data_dir")
             if _CFG.is_file() else None)
XLSX = Path(_DATA_DIR or r"D:\QQBot_Data") / "磁力链接汇总.xlsx"
SHEET = "磁力链接(去重)"
RETRIES = 5


def dedupe() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    hdr = [c.value for c in ws[1]]
    hi = [i for i, v in enumerate(hdr) if v and "特征码" in str(v)][0]
    seen: set[str] = set()
    dup_idx: list[int] = []
    for row in ws.iter_rows(min_row=2):
        h = str(row[hi].value or "").strip().lower()
        if h and h in seen:
            dup_idx.append(row[0].row)
        elif h:
            seen.add(h)
    if not dup_idx:
        wb.close()
        print(f"无重复，共 {len(seen)} 个特征码，不动表")
        return 0
    # 从后往前删：delete_rows 会让后面行的行号前移
    for i in sorted(dup_idx, reverse=True):
        ws.delete_rows(i)
    tmp = XLSX.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    os.replace(tmp, XLSX)   # 直写原文件中断会撕成坏 zip
    wb.close()
    print(f"已删 {len(dup_idx)} 行重复，剩 {len(seen)} 个特征码 -> {XLSX}")
    return len(dup_idx)


def main() -> None:
    if not XLSX.is_file():
        sys.exit(f"登记表不存在: {XLSX}")
    for attempt in range(1, RETRIES + 1):
        try:
            dedupe()
            return
        except PermissionError as e:
            if attempt == RETRIES:
                sys.exit(f"文件被占用（Excel/WPS 开着？）：{e}")
            print(f"文件被占用，5 秒后重试（{attempt}/{RETRIES}）...")
            time.sleep(5)


if __name__ == "__main__":
    main()
